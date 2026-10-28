"""Load the FinAssist investment workbook into PostgreSQL (tables mf_base and mf_txns).

This is a ONE-TIME initial load. After it, the database is the source of truth:
- The database creates all IDs (fund_id, txn_id). IDs from Excel are never stored.
- The loader refuses to run if the tables already contain data, so running it twice
  can never create duplicate holdings or transactions.
"""
import argparse
import os
import re
import sys
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

import openpyxl
import psycopg

# Columns we need from each sheet. Found by NAME, so moving a column in Excel doesn't break us.
# MFBase "ID" and MFTxns "FundId" are read ONLY to link transactions to their holding during
# the load. They are not saved: the database gives every row its own new ID.
BASE_COLUMNS = ["ID", "FolioNo", "SchemeName", "ISIN", "SchemeCode", "SchemeCategory",
                "FirstHolder", "EntityName", "TotalUnitsHeld"]
TXN_COLUMNS = ["FundId", "TransactionDate", "TransactionType", "RealizedAmount",
               "TransactionAmount", "TransactionNAV", "TransactionUnits", "TransactionSTT",
               "TransactionTDS", "TransactionStampDuty"]

TXN_TYPES = {"PUR", "RED", "REI", "DIV", "REV"}
ISIN_PATTERN = re.compile(r"^IN[A-Z0-9]{10}$")   # Indian ISINs: IN + 10 letters/digits


class RowError(ValueError):
    """A value in one row is missing or invalid, so that row is rejected."""


# ---------------------------- read tables -------------------------------------
def read_sheet(workbook, name, required):
    """Return a list of (excel_row_number, {column_name: value}) for every non-empty row."""
    if name not in workbook.sheetnames:
        raise SystemExit(f"Sheet '{name}' not found. Sheets in file: {workbook.sheetnames}")

    rows = workbook[name].iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else None for h in next(rows)]

    missing = [c for c in required if c not in header]
    if missing:
        raise SystemExit(f"Sheet '{name}' is missing columns: {missing}")

    records = []
    for excel_row, values in enumerate(rows, start=2):          # row 1 is the header
        record = {h: v for h, v in zip(header, values) if h is not None}
        if any(v not in (None, "") for v in record.values()):   # skip completely empty rows
            records.append((excel_row, record))
    return records


# ---------------------------- cleaning helpers -------------------------------------
def to_text(value, column, max_len, required=True):
    """Trim text; numbers become text without '.0'; empty -> error or None."""
    if value is None or str(value).strip() == "":
        if required:
            raise RowError(f"{column} is empty")
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)                          # 1016126706.0 -> 1016126706
    text = " ".join(str(value).split())             # trim and collapse double spaces
    if len(text) > max_len:
        raise RowError(f"{column} is longer than {max_len} characters")
    return text


def to_int(value, column, required=True):
    """Whole number; empty -> error, or None if the column is optional."""
    if value is None or str(value).strip() == "":
        if required:
            raise RowError(f"{column} is empty")
        return None
    try:
        number = Decimal(str(value))
    except InvalidOperation:
        raise RowError(f"{column} is not a number: {value!r}") from None
    if number != number.to_integral_value():
        raise RowError(f"{column} is not a whole number: {value!r}")
    return int(number)


def to_decimal(value, column, places):
    """Excel gives floats (inexact). Convert to an exact Decimal with fixed decimal places."""
    if value is None or value == "":
        return Decimal(0)
    try:
        return Decimal(str(value)).quantize(Decimal(1).scaleb(-places), rounding=ROUND_HALF_UP)
    except InvalidOperation:
        raise RowError(f"{column} is not a number: {value!r}") from None


def to_date(value, column):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return datetime.strptime(str(value).strip()[:10], "%Y-%m-%d").date()
    except ValueError:
        raise RowError(f"{column} is not a date: {value!r}") from None


# ---------------------------- validation -------------------------------------
def clean_base(rows, rejected, warnings):
    """Turn MFBase rows into clean dicts for table mf_base. Bad rows go to `rejected`.

    Each dict keeps the Excel ID as "excel_id": it is used to link transactions to
    this holding during the load, and is NOT written to the database.
    """
    clean, seen_ids, seen_holdings = [], set(), {}
    for excel_row, r in rows:
        try:
            excel_id = to_int(r["ID"], "ID")
            isin = to_text(r["ISIN"], "ISIN", 12).upper()
            if not ISIN_PATTERN.match(isin):
                raise RowError(f"ISIN {isin!r} is not a valid Indian ISIN")

            category = to_text(r["SchemeCategory"], "SchemeCategory", 100, required=False)
            if category and not any(ch.isalpha() for ch in category):
                warnings.append(f"MFBase row {excel_row}: "
                                f"SchemeCategory {category!r} is not a category -> stored as empty")
                category = None

            record = {
                "excel_id": excel_id,
                "entity_name": to_text(r["EntityName"], "EntityName", 120),
                "scheme_code": to_int(r["SchemeCode"], "SchemeCode", required=False),
                "scheme_name": to_text(r["SchemeName"], "SchemeName", 200),
                "isin": isin,
                "scheme_category": category,
                "folio_no": to_text(r["FolioNo"], "FolioNo", 30),
                "first_holder": to_text(r["FirstHolder"], "FirstHolder", 50),
            }
            # Same key as the table's UNIQUE constraint: one holding = ISIN + folio + holder.
            holding = (record["isin"], record["folio_no"], record["first_holder"])
            if excel_id in seen_ids:
                raise RowError(f"duplicate ID {excel_id}")
            if holding in seen_holdings:
                raise RowError(f"same ISIN, folio and holder as row {seen_holdings[holding]}")
        except RowError as error:
            rejected.append(f"MFBase row {excel_row}: {error}")
            continue
        seen_ids.add(excel_id)
        seen_holdings[holding] = excel_row
        clean.append(record)
    return clean


def clean_txns(rows, valid_excel_ids, rejected):
    """Turn MFTxns rows into clean dicts for table mf_txns. Bad rows go to `rejected`.

    "excel_fund_id" is the holding's Excel ID. It is swapped for the holding's new
    database ID when writing (see write_to_database).
    """
    clean, today = [], date.today()
    for excel_row, r in rows:
        try:
            record = {
                "excel_fund_id": to_int(r["FundId"], "FundId"),
                "txn_date": to_date(r["TransactionDate"], "TransactionDate"),
                "txn_type": to_text(r["TransactionType"], "TransactionType", 3).upper(),
                "amount": to_decimal(r["TransactionAmount"], "TransactionAmount", 2),
                "realized_amt": to_decimal(r["RealizedAmount"], "RealizedAmount", 2),
                "nav": to_decimal(r["TransactionNAV"], "TransactionNAV", 4),
                "units": to_decimal(r["TransactionUnits"], "TransactionUnits", 4),
                "stt": to_decimal(r["TransactionSTT"], "TransactionSTT", 2),
                "tds": to_decimal(r["TransactionTDS"], "TransactionTDS", 2),
                "stamp_duty": to_decimal(r["TransactionStampDuty"], "TransactionStampDuty", 2),
            }
            if record["excel_fund_id"] not in valid_excel_ids:
                raise RowError(f"FundId {record['excel_fund_id']} has no valid row in MFBase")
            if record["txn_type"] not in TXN_TYPES:
                raise RowError(f"unknown TransactionType {record['txn_type']!r}")
            if record["txn_date"] > today:
                raise RowError(f"TransactionDate {record['txn_date']} is in the future")
            if record["units"] < 0 or record["nav"] < 0:
                raise RowError("units or NAV is negative")
            if record["txn_type"] != "DIV" and (record["units"] == 0 or record["nav"] == 0):
                raise RowError(f"{record['txn_type']} needs units and NAV greater than 0")
        except RowError as error:
            rejected.append(f"MFTxns row {excel_row}: {error}")
            continue
        clean.append(record)
    return clean


# ---------------------------- write to database -------------------------------------
# No ID in the column list: the database creates fund_id and txn_id itself (IDENTITY).
# RETURNING hands back the new fund_id, so transactions can be linked to it.
INSERT_BASE = """
    INSERT INTO mf_base (entity_name, scheme_code, scheme_name, isin,
                         scheme_category, folio_no, first_holder)
    VALUES (%(entity_name)s, %(scheme_code)s, %(scheme_name)s, %(isin)s,
            %(scheme_category)s, %(folio_no)s, %(first_holder)s)
    RETURNING fund_id
"""

INSERT_TXN = """
    INSERT INTO mf_txns (fund_id, txn_date, txn_type, amount, realized_amt,
                         nav, units, stt, tds, stamp_duty)
    VALUES (%(fund_id)s, %(txn_date)s, %(txn_type)s, %(amount)s, %(realized_amt)s,
            %(nav)s, %(units)s, %(stt)s, %(tds)s, %(stamp_duty)s)
"""


class TablesNotEmpty(Exception):
    """The one-time load has already been done."""


def write_to_database(base, txns):
    """Write all rows in ONE transaction: everything is saved, or nothing is."""
    if not os.getenv("PGPASSWORD"):
        raise SystemExit("PGPASSWORD is not set. Run the $env:PGPASSWORD line from Step 3.6 first.")

    # Connection details: environment variables if set, otherwise the port-forward defaults.
    # The password is read by psycopg itself from PGPASSWORD; it is never in this file.
    with psycopg.connect(
        host=os.getenv("PGHOST", "localhost"),
        port=os.getenv("PGPORT", "15432"),
        dbname=os.getenv("PGDATABASE", "portfolio_db"),
        user=os.getenv("PGUSER", "postgres"),
        connect_timeout=10,
    ) as conn:
        with conn.cursor() as cur:
            # Lock both tables until we commit, so two loaders started at the same time
            # cannot both see "empty" and both load. Normal reads still work meanwhile.
            cur.execute("LOCK TABLE mf_base, mf_txns IN EXCLUSIVE MODE")
            already_loaded = cur.execute(
                "SELECT EXISTS (SELECT 1 FROM mf_base) OR EXISTS (SELECT 1 FROM mf_txns)"
            ).fetchone()[0]
            if already_loaded:
                raise TablesNotEmpty()

            # Parents first (the foreign key needs them). Remember Excel ID -> new database ID.
            new_fund_id = {}
            for holding in base:
                new_fund_id[holding["excel_id"]] = cur.execute(INSERT_BASE, holding).fetchone()[0]

            # Children: swap each transaction's Excel FundId for the holding's new fund_id.
            cur.executemany(INSERT_TXN, [{**t, "fund_id": new_fund_id[t["excel_fund_id"]]}
                                         for t in txns])

            holdings = cur.execute("SELECT count(*) FROM mf_base").fetchone()[0]
            transactions = cur.execute("SELECT count(*) FROM mf_txns").fetchone()[0]
    # Leaving the "with psycopg.connect" block COMMITS. Any error inside it ROLLS BACK everything.
    print(f"Database now holds {holdings} holdings and {transactions} transactions.")


# ------------------------- main -------------------------------------
def main():
    parser = argparse.ArgumentParser(description="Load the investment workbook (one-time initial load)")
    parser.add_argument("--file", required=True, type=Path, help="path to the .xlsx file")
    parser.add_argument("--dry-run", action="store_true", help="validate only, write nothing")
    args = parser.parse_args()

    workbook = openpyxl.load_workbook(args.file, read_only=True, data_only=True)
    rejected, warnings = [], []

    base = clean_base(read_sheet(workbook, "MFBase", BASE_COLUMNS), rejected, warnings)
    valid_excel_ids = {b["excel_id"] for b in base}
    txns = clean_txns(read_sheet(workbook, "MFTxns", TXN_COLUMNS), valid_excel_ids, rejected)

    print(f"Valid: {len(base)} holdings, {len(txns)} transactions. "
          f"Rejected: {len(rejected)}. Warnings: {len(warnings)}.")
    for line in rejected:
        print("  REJECTED", line)
    for line in warnings:
        print("  WARNING ", line)

    if not base:
        print("Nothing valid to load.")
        return 2
    if args.dry_run:
        print("Dry run: the database was not touched.")
    else:
        try:
            write_to_database(base, txns)
        except TablesNotEmpty:
            print("Load REFUSED: the tables already contain data. "
                  "The Excel import is a one-time initial load; nothing was changed.")
            return 2
        except psycopg.Error as error:
            print(f"Load FAILED, nothing was changed in the database: {error}")
            return 2
    return 1 if rejected else 0      # exit code: 0 = all good, 1 = some rows rejected, 2 = failed


if __name__ == "__main__":
    sys.exit(main())
